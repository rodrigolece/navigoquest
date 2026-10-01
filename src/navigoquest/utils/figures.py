"""Export paper figures as vector graphics together with their source data.

Everything is controlled by ``notebooks/plot_config.toml``, so a notebook only needs
one line per figure::

    from navigoquest.utils.figures import save_figure

    save_figure(fig, "voc_roc")

Each call writes

* ``<dir>/<name>.<fmt>`` for every format listed in ``formats`` (PDF and SVG), and
* ``<dir>/source_data/<name>.xlsx``, the numbers drawn in the figure, with one tab per
  panel.

A panel is one grid of subplots (one matplotlib ``GridSpec``), matching the lettered
panels of the paper's figures: Figure 4 is built from three grids and gets tabs
``Fig4a``, ``Fig4b`` and ``Fig4c``. Within a tab, the subplots of the panel are stacked
in one table, with a ``subplot`` column saying which subplot each row belongs to.

``build_source_data_workbook()`` (run by ``11_source_data.ipynb``) gathers the tabs of
all figures into a single Excel file.

The data of a subplot (a matplotlib ``Axes``) comes from, in this order:

1. tables attached to the axes by a plotting helper with ``attach_source_data``;
2. automatic extraction of what is drawn on the axes (lines, scatter points, bars,
   error bars, filled bands, polygons, arrows and images).

Tables passed with ``save_figure(fig, name, data=...)`` get tabs of their own.
"""

from __future__ import annotations

import math
import pathlib
import re
import tomllib
import warnings
from collections.abc import Mapping
from string import ascii_lowercase

import matplotlib as mpl
import numpy as np
import pandas as pd
from matplotlib.collections import LineCollection, PathCollection, PolyCollection, QuadMesh
from matplotlib.container import BarContainer, ErrorbarContainer
from matplotlib.patches import FancyArrow, Polygon, Rectangle


CONFIG_FILENAME = "plot_config.toml"
SOURCE_DATA_SUBDIR = "source_data"
EXCEL_MAX_ROWS = 1_048_576

_ATTACHED_ATTR = "_navigoquest_source_data"

DEFAULT_CONFIG: dict[str, dict] = {
    "output": {
        "enabled": True,
        "dir": "../figures",
        "formats": ["pdf", "svg"],
        "bbox_inches": "tight",
        "dpi": 300,
        "transparent": False,
    },
    "source_data": {
        "enabled": True,
        "auto_extract": True,
        "workbook": "SourceData.xlsx",
    },
    "names": {},
}

_DRAWN_NOTE = (
    "Values read from the plotted marks. element: type of mark; series: mark number "
    "within the subplot; label: its legend label; point: index within the mark; "
    "x, y: data coordinates (polar subplots: x = angle in radians, y = radius)."
)


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------


def find_config(start: str | pathlib.Path | None = None) -> pathlib.Path | None:
    """Look for ``plot_config.toml`` in ``start`` (default: cwd), its parents, or their
    ``notebooks`` subfolder."""
    start = pathlib.Path(start or pathlib.Path.cwd()).resolve()
    for folder in (start, *start.parents):
        for candidate in (folder / CONFIG_FILENAME, folder / "notebooks" / CONFIG_FILENAME):
            if candidate.is_file():
                return candidate
    return None


def load_config(path: str | pathlib.Path | None = None) -> dict:
    """Read the plot configuration, filling in defaults for anything left out.

    The output ``dir`` is resolved relative to the folder containing the config file.
    The file is re-read on every call, so edits apply without restarting the kernel.
    """
    config_path = pathlib.Path(path) if path is not None else find_config()
    cfg = {section: dict(values) for section, values in DEFAULT_CONFIG.items()}
    base = pathlib.Path.cwd()

    if config_path is None:
        warnings.warn(f"No {CONFIG_FILENAME} found; using default settings.", stacklevel=2)
    else:
        with open(config_path, "rb") as f:
            user_cfg = tomllib.load(f)
        for section, values in user_cfg.items():
            if not isinstance(values, dict):
                raise TypeError(f"Unexpected top-level key {section!r} in {config_path}")
            cfg.setdefault(section, {}).update(values)
        base = config_path.parent

    out_dir = pathlib.Path(cfg["output"]["dir"]).expanduser()
    cfg["output"]["dir"] = (out_dir if out_dir.is_absolute() else base / out_dir).resolve()
    return cfg


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def attach_source_data(ax, data, name: str | None = None) -> None:
    """Record the data drawn on ``ax``.

    ``save_figure`` exports attached tables for this axes instead of extracting the
    drawn artists automatically. Tables attached under different names are shown as
    separate blocks within the panel's tab.
    """
    tables = getattr(ax, _ATTACHED_ATTR, None)
    if tables is None:
        tables = []
        setattr(ax, _ATTACHED_ATTR, tables)
    tables.append((name, _to_frame(data)))


def save_figure(fig, name: str, data=None, auto: bool | None = None, config: dict | None = None):
    """Save ``fig`` in every configured format and export its source data.

    Args:
        fig:    Matplotlib figure.
        name:   Key of the figure. ``[names]`` in the config can map it to another
                output name (e.g. ``voc_roc = "Fig4"``).
        data:   Optional extra table(s), each exported to a tab of its own: a DataFrame,
                Series, array, or a dict mapping tab names to any of those.
        auto:   Whether to extract the drawn data of axes that have no attached data.
                Defaults to ``source_data.auto_extract`` in the config. Pass ``False``
                when ``data`` already covers the whole figure.
        config: Configuration dict (default: read ``plot_config.toml``).

    The files written are listed in the printed message.
    """
    cfg = config if config is not None else load_config()
    out_cfg = cfg["output"]
    if not out_cfg.get("enabled", True):
        return

    out_name = str(cfg.get("names", {}).get(name) or name)
    if not out_name or re.search(r"[\\/:*?\"<>|\[\]]", out_name):
        raise ValueError(f"Invalid figure name {out_name!r}")

    out_dir: pathlib.Path = out_cfg["dir"]
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[str] = []

    save_kwargs = {
        "bbox_inches": out_cfg.get("bbox_inches") or None,
        "dpi": out_cfg.get("dpi", 300),
        "transparent": out_cfg.get("transparent", False),
    }
    # Keep text as real (editable) text: TrueType fonts embedded in PDF, <text> in SVG
    vector_rc = {"pdf.fonttype": 42, "ps.fonttype": 42, "svg.fonttype": "none"}
    with mpl.rc_context(vector_rc):
        for fmt in out_cfg.get("formats", []):
            fmt = str(fmt).lower().lstrip(".")
            path = out_dir / f"{out_name}.{fmt}"
            fig.savefig(path, format=fmt, **save_kwargs)
            written.append(path.name)

    sd_cfg = cfg.get("source_data", {})
    if sd_cfg.get("enabled", True):
        if auto is None:
            auto = sd_cfg.get("auto_extract", True)
        tabs = collect_source_data(fig, name=out_name, data=data, auto=auto)
        if tabs:
            path = out_dir / SOURCE_DATA_SUBDIR / f"{out_name}.xlsx"
            path.parent.mkdir(parents=True, exist_ok=True)
            sheet_names = _write_workbook(path, tabs)
            written.append(f"{SOURCE_DATA_SUBDIR}/{path.name} (tabs: {', '.join(sheet_names)})")

    print(f"Saved '{out_name}': " + ", ".join(written))


def collect_source_data(fig, name: str = "Figure", data=None, auto: bool = True) -> list:
    """Group the source data of a figure into tabs.

    Returns ``[(tab_name, description, [(block_title, DataFrame), ...]), ...]``: one
    tab per panel (grid of subplots) that has data, then one per explicit table.
    """
    groups: dict = {}
    headings: dict = {}  # titles of data-less axes spanning a grid (e.g. a panel heading)
    for ax in fig.axes:
        if not ax.get_visible() or getattr(ax, "_colorbar", None) is not None:
            continue
        ss = _subplotspec(ax)
        key = id(ss.get_gridspec()) if ss is not None else None
        tables = _axes_tables(ax, auto)
        if tables:
            groups.setdefault(key, []).append((ax, tables))
        elif ax.get_title() and key is not None:
            headings.setdefault(key, []).append(_plain(ax.get_title()))

    tabs = []
    prefix = name if re.search(r"\d$", name) else f"{name}_"
    for letter, (key, entries) in zip(ascii_lowercase, groups.items(), strict=False):
        tab = f"{prefix}{letter}" if len(groups) > 1 else name
        axes = [ax for ax, _ in entries]
        labels = _subplot_labels(axes)

        blocks: dict[str, list[pd.DataFrame]] = {}
        for (_ax, tables), label in zip(entries, labels, strict=True):
            for kind, df in tables:
                df = _flat(df)
                if len(entries) > 1:
                    df.insert(0, "subplot", label)
                blocks.setdefault(kind, []).append(df)

        block_list = [
            (_block_title(kind), pd.concat(dfs, ignore_index=True)) for kind, dfs in blocks.items()
        ]
        description = _panel_description(axes, labels)
        if headings.get(key):
            description = f"{'; '.join(headings[key])}. {description}"
        tabs.append((tab, description, block_list))

    for key, df in _explicit_tables(data):
        tab = name if key is None else f"{name}_{key}"
        tabs.append((tab, "Data underlying the figure.", [("", _flat(df))]))

    return tabs


def extract_axes_data(ax) -> pd.DataFrame | None:
    """Extract the numbers drawn on an axes as a long-format table.

    Columns: ``element`` (kind of artist), ``series`` (artist number), ``label`` (its
    legend label, if any), ``point`` (index within the series), ``x``, ``y`` and, where
    relevant, extra columns such as ``value`` (colour value), ``y_lower``/``y_upper``
    (bands and error bars), bar geometry, or arrow components ``dx``/``dy``.
    For polar axes ``x`` is the angle in radians and ``y`` the radius.
    """
    frames: list[pd.DataFrame] = []
    consumed: set[int] = set()

    def add(element: str, label, df: pd.DataFrame) -> None:
        if df is None or df.empty:
            return
        df = df.reset_index(drop=True)
        df.insert(0, "point", np.arange(len(df)))
        df.insert(0, "label", _clean_label(label))
        df.insert(0, "series", len(frames) + 1)
        df.insert(0, "element", element)
        frames.append(df)

    # Containers first, so that their parts are not reported twice
    for container in ax.containers:
        if isinstance(container, ErrorbarContainer):
            data_line, caplines, barlinecols = container.lines
            parts = [data_line, *caplines, *barlinecols]
            consumed.update(id(a) for a in parts if a is not None)
            add("errorbar", container.get_label(), _errorbar_frame(data_line, barlinecols))
        elif isinstance(container, BarContainer):
            consumed.update(id(p) for p in container.patches)
            horizontal = getattr(container, "orientation", "vertical") == "horizontal"
            rows = []
            for r in container.patches:
                x0, y0, w, h = r.get_x(), r.get_y(), r.get_width(), r.get_height()
                # x, y: bar position and the value it shows (the end of the bar)
                x, y = (x0 + w, y0 + h / 2) if horizontal else (x0 + w / 2, y0 + h)
                rows.append({"x": x, "y": y, "x_left": x0, "width": w, "y_bottom": y0, "height": h})
            add("barh" if horizontal else "bar", container.get_label(), pd.DataFrame(rows))

    for line in ax.get_lines():
        if id(line) in consumed or not line.get_visible():
            continue
        xy = _float_array(line.get_xydata()).reshape(-1, 2)
        if len(xy):
            add("line", line.get_label(), pd.DataFrame(xy, columns=["x", "y"]))

    for coll in ax.collections:
        if id(coll) in consumed or not coll.get_visible() or coll.get_alpha() == 0:
            continue
        label = coll.get_label()
        if isinstance(coll, PathCollection):
            offsets = _float_array(coll.get_offsets()).reshape(-1, 2)
            df = pd.DataFrame(offsets, columns=["x", "y"])
            values = coll.get_array()
            if values is not None and np.size(values) == len(df):
                df["value"] = _float_array(values).ravel()
            add("scatter", label, df)
        elif type(coll).__name__ == "FillBetweenPolyCollection":  # from fill_between
            for path in coll.get_paths():
                add("band", label, _band_frame(path.vertices, getattr(coll, "t_direction", "x")))
        elif isinstance(coll, PolyCollection):
            for path in coll.get_paths():
                add("polygon", label, pd.DataFrame(_float_array(path.vertices), columns=["x", "y"]))
        elif isinstance(coll, LineCollection):
            for seg in coll.get_segments():
                add("segment", label, pd.DataFrame(_float_array(seg), columns=["x", "y"]))
        elif isinstance(coll, QuadMesh):
            add("mesh", label, _matrix_frame(coll.get_array()))

    arrows = []
    for patch in ax.patches:
        if id(patch) in consumed or not patch.get_visible():
            continue
        if isinstance(patch, FancyArrow) and hasattr(patch, "_dx"):
            arrows.append({"x": patch._x, "y": patch._y, "dx": patch._dx, "dy": patch._dy})
        elif isinstance(patch, Polygon):
            xy = _float_array(patch.get_xy())
            add("polygon", patch.get_label(), pd.DataFrame(xy, columns=["x", "y"]))
        elif isinstance(patch, Rectangle):
            row = {
                "x": patch.get_x(),
                "y": patch.get_y(),
                "width": patch.get_width(),
                "height": patch.get_height(),
            }
            add("rectangle", patch.get_label(), pd.DataFrame([row]))
    if arrows:
        add("arrow", None, pd.DataFrame(arrows))

    for image in ax.get_images():
        arr = np.asarray(image.get_array())
        if arr.ndim == 2:
            add("image", image.get_label(), _matrix_frame(arr))

    if not frames:
        return None
    return pd.concat(frames, ignore_index=True)


def build_source_data_workbook(config: dict | None = None, path: str | pathlib.Path | None = None):
    """Combine the source data of all saved figures into one Excel workbook.

    Copies every tab of the per-figure files in ``<dir>/source_data`` (main figures
    first, then supplementary ones, then the rest), adds any CSV files found there
    (e.g. written by the R script) as tabs, and starts with a Contents tab.
    Returns the path written.
    """
    from openpyxl import Workbook, load_workbook

    cfg = config if config is not None else load_config()
    out_dir: pathlib.Path = cfg["output"]["dir"]
    root = out_dir / SOURCE_DATA_SUBDIR
    if path is None:
        path = out_dir / cfg.get("source_data", {}).get("workbook", "SourceData.xlsx")
    path = pathlib.Path(path)

    files = [
        p for p in root.glob("*") if p.suffix in (".xlsx", ".csv") and not p.name.startswith("~$")
    ]
    files = sorted(files, key=_sheet_order) if root.is_dir() else []
    if not files:
        raise FileNotFoundError(f"No source data found in {root}; run the notebooks first.")

    wb = Workbook(write_only=True)
    contents = wb.create_sheet("Contents")
    used = {"contents"}
    toc = []

    for file in files:
        if file.suffix == ".csv":
            df = pd.read_csv(file)
            title = _write_tab(wb, file.stem, f"From {file.name}.", [("", df)], used)[0]
            toc.append((title, file.name, f"From {file.name}."))
            continue
        src = load_workbook(file, read_only=True)
        for ws_src in src.worksheets:
            ws = wb.create_sheet(_sheet_name(ws_src.title, used))
            _copy_widths(ws_src, ws)
            description = ""
            for i, row in enumerate(ws_src.iter_rows()):
                cells = [c for c in row]
                if i == 1 and cells:
                    description = cells[0].value or ""
                font = cells[0].font if cells and cells[0].value is not None else None
                if font is not None and (font.b or font.i):
                    ws.append([_styled(ws, c.value, font) for c in cells])
                else:
                    ws.append([c.value for c in cells])
            toc.append((ws.title, file.name, description))
        src.close()

    _append_styled(contents, ["Tab", "File", "Description"], bold=True)
    for row in toc:
        contents.append(list(row))
    _set_widths(contents, [18, 30, 100])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"Wrote {path} ({len(toc)} tab(s) of source data)")
    return path


# ---------------------------------------------------------------------------
# Helpers: collecting tables
# ---------------------------------------------------------------------------


def _to_frame(data) -> pd.DataFrame:
    if isinstance(data, pd.DataFrame):
        return data.copy()
    if isinstance(data, pd.Series):
        return data.to_frame()
    return pd.DataFrame(data)


def _flat(df: pd.DataFrame) -> pd.DataFrame:
    """Turn a meaningful index into ordinary columns."""
    df = df.copy()
    if isinstance(df.index, pd.RangeIndex):
        return df
    unnamed = [n is None for n in df.index.names]
    df = df.reset_index()
    if len(unnamed) == 1 and unnamed[0]:
        df = df.rename(columns={"index": ""})
    return df


def _explicit_tables(data) -> list[tuple[str | None, pd.DataFrame]]:
    if data is None:
        return []
    if isinstance(data, Mapping):
        return [(_slug(str(k)), _to_frame(v)) for k, v in data.items()]
    return [(None, _to_frame(data))]


def _axes_tables(ax, auto: bool) -> list[tuple[str, pd.DataFrame]]:
    attached = getattr(ax, _ATTACHED_ATTR, None)
    if attached:
        return [(n or "", df) for n, df in attached]
    if auto:
        df = extract_axes_data(ax)
        if df is not None and not df.empty:
            return [("\0drawn", df)]
    return []


def _block_title(kind: str) -> str:
    if kind == "\0drawn":
        return _DRAWN_NOTE
    return kind.replace("_", " ").capitalize()


def _subplotspec(ax):
    ss = ax.get_subplotspec() if hasattr(ax, "get_subplotspec") else None
    return ss.get_topmost_subplotspec() if ss is not None else None


def _plain(text: str) -> str:
    """Render simple matplotlib mathtext as plain text (e.g. $\\epsilon3$ -> ε3)."""
    text = str(text)
    for tex, char in (("\\epsilon", "ε"), ("\\alpha", "α"), ("\\beta", "β"), ("\\mu", "μ")):
        text = text.replace(tex, char)
    return re.sub(r"[${}\\]", "", text).strip()


def _subplot_labels(axes) -> list[str]:
    titles = [_plain(ax.get_title() or ax.get_title("left")) for ax in axes]
    if all(titles) and len(set(titles)) == len(titles):
        return titles
    labels = []
    for i, (ax, title) in enumerate(zip(axes, titles, strict=True)):
        ss = _subplotspec(ax)
        pos = f"row {ss.rowspan.start + 1}, col {ss.colspan.start + 1}" if ss else f"axes {i + 1}"
        labels.append(f"{pos} ({title})" if title else pos)
    return labels


def _panel_description(axes, labels) -> str:
    parts = []
    if len(axes) > 1:
        parts.append(f"{len(axes)} subplots: {', '.join(labels)}.")
    elif labels[0]:
        parts.append(f"Subplot: {labels[0]}.")
    for which, getter in (("x axis", "get_xlabel"), ("y axis", "get_ylabel")):
        names = []
        for ax in axes:
            text = _plain(getattr(ax, getter)())
            if text and text not in names:
                names.append(text)
        if names:
            parts.append(f"{which}: {', '.join(names)}.")
    return " ".join(parts)


# ---------------------------------------------------------------------------
# Helpers: writing Excel
# ---------------------------------------------------------------------------


def _write_workbook(path: pathlib.Path, tabs) -> list[str]:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    used: set[str] = set()
    names = []
    for tab, description, blocks in tabs:
        names.extend(_write_tab(wb, tab, description, blocks, used))
    wb.save(path)
    return names


def _write_tab(wb, tab: str, description: str, blocks, used: set[str]) -> list[str]:
    """Write a titled tab; a block that would overflow Excel's row limit continues on
    a new tab. Returns the tab names created."""
    created = []

    def new_sheet(part: int):
        title = tab if part == 1 else f"{tab} ({part})"
        ws = wb.create_sheet(_sheet_name(title, used))
        _set_widths(ws, [16] * max(len(df.columns) for _, df in blocks) if blocks else [])
        _append_styled(ws, [ws.title], bold=True)
        _append_styled(ws, [description], italic=True)
        created.append(ws.title)
        return ws, 2

    ws, n_rows = new_sheet(1)
    for block_title, df in blocks:
        header = ["" if str(c).startswith("Unnamed: ") else str(c) for c in df.columns]
        rows = df.itertuples(index=False, name=None)
        start = True
        for row in rows:
            if start or n_rows >= EXCEL_MAX_ROWS - 1:
                if not start:
                    ws, n_rows = new_sheet(len(created) + 1)
                ws.append([])
                if block_title:
                    _append_styled(ws, [block_title], bold=True)
                    n_rows += 1
                _append_styled(ws, header, bold=True)
                n_rows += 2
                start = False
            ws.append([_excel_value(v) for v in row])
            n_rows += 1
    return created


def _styled(ws, value, font):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    cell = WriteOnlyCell(ws, value=value)
    cell.font = Font(bold=font.b, italic=font.i)
    return cell


def _append_styled(ws, values, bold: bool = False, italic: bool = False) -> None:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    cells = []
    for v in values:
        cell = WriteOnlyCell(ws, value=_excel_value(v))
        cell.font = Font(bold=bold, italic=italic)
        cells.append(cell)
    ws.append(cells)


def _set_widths(ws, widths) -> None:
    from openpyxl.utils import get_column_letter

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _copy_widths(src, dst) -> None:
    try:
        dims = src.column_dimensions
    except AttributeError:  # read-only worksheets do not expose column widths
        dims = {}
    for letter, dim in dims.items():
        if dim.width:
            dst.column_dimensions[letter].width = dim.width


# ---------------------------------------------------------------------------
# Helpers: extraction and names
# ---------------------------------------------------------------------------


def _clean_label(label) -> str:
    label = "" if label is None else str(label)
    return "" if label.startswith("_") else label


def _float_array(values) -> np.ndarray:
    return np.ma.filled(np.ma.asarray(values, dtype=float), np.nan)


def _errorbar_frame(data_line, barlinecols) -> pd.DataFrame:
    df = pd.DataFrame()
    if data_line is not None:
        xy = _float_array(data_line.get_xydata()).reshape(-1, 2)
        df = pd.DataFrame(xy, columns=["x", "y"])
    for lc in barlinecols:
        segs = [_float_array(s) for s in lc.get_segments()]
        if not segs or any(s.shape != (2, 2) for s in segs):
            continue
        segs = np.stack(segs)
        vertical = np.allclose(segs[:, 0, 0], segs[:, 1, 0], equal_nan=True)
        axis, prefix = (1, "y") if vertical else (0, "x")
        lower, upper = segs[:, :, axis].min(axis=1), segs[:, :, axis].max(axis=1)
        if df.empty:
            other = 0 if vertical else 1
            df = pd.DataFrame({"x" if vertical else "y": segs[:, 0, other]})
        if len(lower) == len(df):
            df[f"{prefix}_lower"] = lower
            df[f"{prefix}_upper"] = upper
    return df


def _band_frame(vertices, direction: str = "x") -> pd.DataFrame:
    v = _float_array(vertices)
    t_col, v_col = (0, 1) if direction == "x" else (1, 0)
    df = pd.DataFrame({"t": v[:, t_col], "v": v[:, v_col]}).dropna()
    grouped = df.groupby("t", sort=True)["v"].agg(["min", "max"]).reset_index()
    if direction == "x":
        grouped.columns = ["x", "y_lower", "y_upper"]
    else:
        grouped.columns = ["y", "x_lower", "x_upper"]
    return grouped


def _matrix_frame(arr) -> pd.DataFrame:
    arr = _float_array(arr)
    if arr.ndim != 2:
        return pd.DataFrame({"value": arr.ravel()})
    rows, cols = np.indices(arr.shape)
    return pd.DataFrame({"row": rows.ravel(), "col": cols.ravel(), "value": arr.ravel()})


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_") or "table"


def _sheet_order(p: pathlib.Path):
    """Main figures first, then supplementary figures, then the rest; numbers in natural order."""
    name = p.name.lower()
    group = 0 if re.match(r"fig", name) else 1 if re.match(r"(supp|sfig|ext)", name) else 2
    return group, [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", name)]


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
    candidate, k = base, 2
    while candidate.lower() in used:
        suffix = f"~{k}"
        candidate = base[: 31 - len(suffix)] + suffix
        k += 1
    used.add(candidate.lower())
    return candidate


def _excel_value(v):
    if v is None:
        return None
    if isinstance(v, (float, np.floating)):
        return None if math.isnan(v) else float(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.bool_):
        return bool(v)
    return v
